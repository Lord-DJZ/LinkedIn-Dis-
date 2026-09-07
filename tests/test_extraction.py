import os
import pymupdf
import docx
import pytest
from app.extraction.pdf import PDFExtractor
from app.extraction.docx import DOCXExtractor
from app.extraction.rules import RuleBasedResumeParser


def test_pdf_extraction(tmp_path):
    pdf_file = tmp_path / "sample_resume.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Jane Doe\nSoftware Engineer\nEmail: jane.doe@techhub.com\nPhone: +1 555-987-6543\nSkills: Python, FastAPI, PostgreSQL, Docker\nEducation: BSc in Computer Science, University of Colombo (2018-2022)")
    doc.save(str(pdf_file))
    doc.close()

    text, is_scanned = PDFExtractor.extract_text(str(pdf_file))
    assert not is_scanned
    assert "Jane Doe" in text
    assert "jane.doe@techhub.com" in text
    assert "FastAPI" in text


def test_scanned_pdf_detection(tmp_path):
    # Empty or very short PDF simulates a pure scanned image
    pdf_file = tmp_path / "scanned_image.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Hi")  # Less than 50 chars threshold
    doc.save(str(pdf_file))
    doc.close()

    text, is_scanned = PDFExtractor.extract_text(str(pdf_file))
    assert is_scanned is True


def test_docx_extraction(tmp_path):
    docx_file = tmp_path / "sample_resume.docx"
    doc = docx.Document()
    doc.add_paragraph("Michael Scott")
    doc.add_paragraph("Senior Backend Engineer")
    doc.add_paragraph("Email: michael@dundermifflin.com | Phone: +1 555 432 1098")
    doc.add_paragraph("GitHub: https://github.com/mscott")
    doc.add_paragraph("Skills: Python, Django, PostgreSQL, Docker, AWS")
    doc.save(str(docx_file))

    text, is_empty = DOCXExtractor.extract_text(str(docx_file))
    assert not is_empty
    assert "Michael Scott" in text
    assert "michael@dundermifflin.com" in text
    assert "Django" in text


def test_rule_based_parser():
    sample_text = """
    Kasun Rajapaksha
    Senior Software Engineer
    Email: kasun.r@example.com
    Phone: +94 77 123 4567
    LinkedIn: https://linkedin.com/in/kasun-rajapaksha
    GitHub: https://github.com/kasunr

    TECHNICAL SKILLS
    Python, FastAPI, PostgreSQL, Docker, React, REST API

    EDUCATION
    BSc in Computer Science, University of Moratuwa (2016 - 2020)

    EXPERIENCE
    Virtusa - Senior Software Engineer (2020 - Present)
    Architected high-scale REST APIs using FastAPI and PostgreSQL.
    """

    res = RuleBasedResumeParser.parse(sample_text)
    assert res.personal_information.email == "kasun.r@example.com"
    assert res.personal_information.phone == "+94 77 123 4567"
    assert any("linkedin.com/in/kasun-rajapaksha" in link for link in res.portfolio_links)
    assert any("github.com/kasunr" in link for link in res.portfolio_links)

    skill_names = [s.normalized_name for s in res.skills]
    assert "Python" in skill_names
    assert "FastAPI" in skill_names
    assert "PostgreSQL" in skill_names

    assert len(res.education) > 0
    assert res.education[0].normalized_degree_level == "bachelor"
